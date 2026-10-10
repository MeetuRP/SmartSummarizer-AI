"""Run reproducible synthetic text, document, and vision cases against Ollama."""
import io
import json
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import pymupdf
from docx import Document
from PIL import Image, ImageDraw

from services.extraction_service import extract_text_from_file
from services.local_ai_service import generate
from services.summarizer_service import generate_structured_summary


def synthetic_image(kind: str) -> bytes:
    image = Image.new("RGB", (900, 560), "white")
    draw = ImageDraw.Draw(image)
    if kind == "text":
        draw.text((45, 45), "NORTHSTAR LIBRARY\nHours change: May 3\nWeekdays: 8 am to 8 pm", fill="black", spacing=18)
    elif kind == "chart":
        draw.text((40, 25), "Fictional monthly book returns", fill="black")
        draw.text((40, 65), "January 20 books", fill="black")
        draw.text((40, 100), "February 30 books", fill="black")
        draw.text((40, 135), "March 25 books", fill="black")
        draw.line((300, 450, 300, 180), fill="black", width=4)
        for x, height, label in [(380, 150, "20"), (520, 220, "30"), (660, 185, "25")]:
            draw.rectangle((x, 450-height, x+75, 450), fill="#278779")
            draw.text((x+20, 430-height), label, fill="black")
        for x, label in [(380, "Jan"), (520, "Feb"), (660, "Mar")]:
            draw.text((x+15, 465), label, fill="black")
    else:
        draw.rectangle((0, 0, 900, 560), fill="#cfe7e7")
        draw.ellipse((650, 40, 790, 180), fill="#e3ad65")
        draw.rectangle((0, 390, 900, 560), fill="#75966b")
        draw.polygon([(160, 390), (310, 160), (460, 390)], fill="#66849b")
        draw.polygon([(390, 390), (570, 130), (760, 390)], fill="#839bad")
    output = io.BytesIO()
    image.save(output, "PNG")
    return output.getvalue()


def run_case(name: str, expected_facts: list[str], operation) -> dict:
    started = time.perf_counter()
    output = operation()
    elapsed = time.perf_counter() - started
    return {
        "case": name,
        "expected_facts_for_manual_review": expected_facts,
        "elapsed_seconds": round(elapsed, 3),
        "output": output,
        "nonempty": bool(output.strip()),
    }


def main() -> int:
    short_text = "Northstar Library starts extended weekday hours on May 3. It will close at 8 pm instead of 6 pm."
    long_text = ("Fictional museum archive: the East Hall renovation begins on June 12 and costs $240,000. "
                 "The gallery stays open during work, but the north entrance is closed. " * 80)

    word = Document()
    word.add_heading("Quarterly update", 0)
    word.add_paragraph("The fictional team delivered 18 of 20 planned tasks in Q2.")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Metric", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "On-time tasks", "18 of 20"
    word_bytes = io.BytesIO()
    word.save(word_bytes)

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), short_text)
    pdf_bytes = pdf.tobytes()
    pdf.close()

    text_image = synthetic_image("text")
    chart_image = synthetic_image("chart")
    scene_image = synthetic_image("scene")

    cases = [
        run_case("short_text", ["May 3", "8 pm", "weekday"], lambda: generate_structured_summary(short_text)),
        run_case("multi_chunk_text", ["June 12", "$240,000", "north entrance closed"], lambda: generate_structured_summary(long_text)),
        run_case("text_pdf", ["May 3", "8 pm"], lambda: generate_structured_summary(extract_text_from_file("sample.pdf", "application/pdf", pdf_bytes)[1])),
        run_case("docx_with_table", ["18 of 20", "Q2"], lambda: generate_structured_summary(extract_text_from_file("sample.docx", None, word_bytes.getvalue())[1])),
        run_case("image_with_text", ["May 3", "8 am", "8 pm"], lambda: generate("Read the visible text; preserve dates and times and mark uncertainty.", image_bytes=text_image)),
        run_case("synthetic_chart", ["January 20 books", "February 30 books", "March 25 books"], lambda: generate("Read the chart labels and values. Do not infer unreadable values.", image_bytes=chart_image)),
        run_case("scene_without_data", ["mountains", "sun", "landscape"], lambda: generate("Describe the scene without inventing text or numerical data.", image_bytes=scene_image)),
    ]

    root = Path(__file__).resolve().parents[1]
    report_dir = root / "artifacts" / "evaluation"
    report_dir.mkdir(parents=True, exist_ok=True)
    report_path = report_dir / "latest.json"
    report = {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "model": __import__("config").get_settings().local_vision_model,
        "type": "real_local_inference_synthetic_eval",
        "manual_review_required": True,
        "cases": cases,
        "audio_fixture": "not included; add a licensed local speech sample to evaluate transcription",
    }
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(f"Wrote {report_path}")
    print(f"Completed {sum(case['nonempty'] for case in cases)}/{len(cases)} nonempty model responses. Review each output against its expected facts.")
    return 0 if all(case["nonempty"] for case in cases) else 1


if __name__ == "__main__":
    raise SystemExit(main())
