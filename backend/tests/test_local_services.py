import io
import unittest
import urllib.error
import wave
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import av
import numpy as np
from PIL import Image
from docx import Document
import pymupdf
from fastapi import HTTPException, UploadFile

from services.extraction_service import validate_image, extract_text_from_file
from services.extraction_service import _extract_video_audio_bytes
from services import local_ai_service
from services.summarizer_service import generate_structured_summary, process_uploaded_file
from routers.summary_router import upload_and_summarize
from utils.chunking import chunk_text


class ChunkTests(unittest.TestCase):
    def test_chunking_overlap_and_empty(self):
        self.assertEqual(chunk_text("", 5), [])
        chunks = chunk_text("abcdefghij", 5, 2)
        self.assertEqual(chunks, ["abcde", "defgh", "ghij"])


class ImageTests(unittest.TestCase):
    def test_image_validation_normalizes_orientation_and_format(self):
        buffer = io.BytesIO()
        Image.new("RGBA", (4, 3), (255, 0, 0, 120)).save(buffer, "PNG")
        normalized = validate_image(buffer.getvalue())
        with Image.open(io.BytesIO(normalized)) as decoded:
            self.assertEqual(decoded.format, "JPEG")
            self.assertEqual(decoded.mode, "RGB")

    def test_corrupt_image_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "invalid"):
            validate_image(b"not an image")

    def test_image_upload_extracts_as_image(self):
        buffer = io.BytesIO()
        Image.new("RGB", (2, 2)).save(buffer, "PNG")
        kind, text, image = extract_text_from_file("chart.png", "image/png", buffer.getvalue())
        self.assertEqual(kind, "image")
        self.assertIn("visual analysis", text)
        self.assertTrue(image)

    def test_txt_and_unsupported_file(self):
        kind, text, image = extract_text_from_file("notes.txt", "text/plain", "Caf\\u00e9".encode("utf-8"))
        self.assertEqual((kind, text, image), ("txt", "Caf\\u00e9", None))
        with self.assertRaisesRegex(ValueError, "Unsupported"):
            extract_text_from_file("payload.bin", "application/octet-stream", b"data")

    def test_pdf_text_and_docx_table_extraction(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Annual report revenue increased by 12 percent. " * 4)
        pdf = document.tobytes()
        document.close()
        pdf_type, pdf_text, _ = extract_text_from_file("report.pdf", "application/pdf", pdf)
        self.assertEqual(pdf_type, "pdf")
        self.assertIn("Annual report", pdf_text)

        word = Document()
        word.add_paragraph("Quarterly update")
        table = word.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Revenue"
        table.cell(0, 1).text = "$42M"
        stream = io.BytesIO()
        word.save(stream)
        doc_type, doc_text, _ = extract_text_from_file("update.docx", None, stream.getvalue())
        self.assertEqual(doc_type, "docx")
        self.assertIn("Revenue | $42M", doc_text)

    @patch("services.local_ai_service.generate", return_value="Scanned page text: Invoice total $42.")
    def test_scanned_pdf_page_uses_local_vision(self, vision):
        image = io.BytesIO()
        Image.new("RGB", (80, 50), "white").save(image, "PNG")
        document = pymupdf.open()
        page = document.new_page()
        page.insert_image(pymupdf.Rect(30, 30, 350, 200), stream=image.getvalue())
        pdf = document.tobytes()
        document.close()
        kind, text, _ = extract_text_from_file("scan.pdf", "application/pdf", pdf)
        self.assertEqual(kind, "pdf")
        self.assertIn("Invoice total $42", text)
        vision.assert_called_once()

    @patch("services.extraction_service._extract_video_audio_bytes", return_value=b"wav bytes")
    @patch("services.extraction_service.transcribe_audio_bytes", return_value="spoken words")
    def test_audio_and_video_use_local_transcription_adapter(self, transcribe, _extract):
        audio_type, audio_text, _ = extract_text_from_file("memo.wav", "audio/wav", b"audio")
        video_type, video_text, _ = extract_text_from_file("meeting.mp4", "video/mp4", b"video")
        self.assertEqual((audio_type, audio_text), ("audio", "spoken words"))
        self.assertEqual((video_type, video_text), ("video", "spoken words"))
        self.assertEqual(transcribe.call_count, 2)

    @unittest.skipUnless("aac" in av.codecs_available, "AAC encoder is unavailable")
    def test_video_audio_is_resampled_to_local_whisper_wav(self):
        video = io.BytesIO()
        container = av.open(video, mode="w", format="mp4")
        stream = container.add_stream("aac", rate=16000)
        stream.layout = "mono"
        for index in range(8):
            frame = av.AudioFrame.from_ndarray(
                np.zeros((1, 1024), dtype=np.int16), format="s16", layout="mono"
            )
            frame.sample_rate = 16000
            frame.pts = index * 1024
            for packet in stream.encode(frame):
                container.mux(packet)
        for packet in stream.encode():
            container.mux(packet)
        container.close()

        wav_bytes = _extract_video_audio_bytes(video.getvalue())
        with wave.open(io.BytesIO(wav_bytes), "rb") as wav_file:
            self.assertEqual(wav_file.getnchannels(), 1)
            self.assertEqual(wav_file.getframerate(), 16000)
            self.assertGreater(wav_file.getnframes(), 0)


class LocalAIRequestTests(unittest.TestCase):
    @patch("services.local_ai_service._ollama_request", return_value={"message": {"content": "A grounded summary."}})
    def test_generation_uses_local_adapter_and_image_payload(self, request):
        result = local_ai_service.generate("Summarize", image_bytes=b"small image")
        self.assertEqual(result, "A grounded summary.")
        args, kwargs = request.call_args
        self.assertEqual(args[0], "/api/chat")
        self.assertIn("images", args[1]["messages"][1])
        self.assertIs(args[1]["think"], False)

    @patch("services.local_ai_service._ollama_request", return_value={"message": {"content": "  "}})
    def test_empty_model_response_is_an_error(self, _request):
        with self.assertRaisesRegex(RuntimeError, "empty response"):
            local_ai_service.generate("prompt")

    def test_remote_ollama_endpoint_is_rejected(self):
        with patch.dict("os.environ", {"OLLAMA_BASE_URL": "https://example.com"}):
            from config import Settings
            with self.assertRaisesRegex(ValueError, "loopback"):
                Settings()

    @patch("services.local_ai_service.urllib.request.urlopen", side_effect=urllib.error.URLError("connection refused"))
    def test_ollama_unavailable_error_is_actionable(self, _urlopen):
        with self.assertRaisesRegex(local_ai_service.LocalAIUnavailableError, "Start Ollama"):
            local_ai_service.list_models()

    def test_upload_maps_local_ai_outage_to_service_unavailable(self):
        async def run_request():
            from unittest.mock import AsyncMock
            with patch(
                "routers.summary_router.process_uploaded_file",
                new=AsyncMock(side_effect=local_ai_service.LocalAIUnavailableError("Start Ollama")),
            ):
                await upload_and_summarize(UploadFile(file=io.BytesIO(b"text"), filename="note.txt"))

        with self.assertRaises(HTTPException) as raised:
            import asyncio
            asyncio.run(run_request())
        self.assertEqual(raised.exception.status_code, 503)

    @patch("services.summarizer_service.generate", return_value="Key Points\\n- A fact\\nImportant Details\\n...\\nFinal Summary\\n...")
    def test_long_document_uses_chunk_summaries_and_merge(self, generate):
        with patch("services.summarizer_service.get_settings") as settings:
            settings.return_value.max_chunk_chars = 100
            settings.return_value.chunk_overlap = 10
            output = generate_structured_summary("Revenue rose. " * 20)
        self.assertIn("Final Summary", output)
        self.assertGreaterEqual(generate.call_count, 2)

    @patch(
        "services.summarizer_service.generate",
        return_value=(
            "Key Points\n- The speaker was John F. Kennedy.\n- This occurred in 1943 during wartime.\n"
            "Important Details\n- Historical context: an inaugural address.\nFinal Summary\n..."
        ),
    )
    def test_unsupported_quote_attribution_uses_verbatim_fallback(self, _generate):
        source = "And so my fellow Americans, ask not what your country can do for you, ask what you can do for your country."
        result = generate_structured_summary(source)
        self.assertIn("unsupported claims were omitted", result)
        self.assertIn(source, result)
        self.assertNotIn("John F. Kennedy", result)
        self.assertNotIn("1943", result)
        self.assertNotIn("wartime", result)

    def test_upload_processing_logs_visible_stages(self):
        collection = SimpleNamespace(
            insert_one=AsyncMock(return_value=SimpleNamespace(inserted_id="summary-id"))
        )

        async def run_upload():
            with (
                patch("services.summarizer_service.extract_text_from_file", return_value=("txt", "Source text.", None)),
                patch("services.summarizer_service.generate_structured_summary", return_value="Summary text."),
                patch("services.summarizer_service.get_summary_collection", return_value=collection),
            ):
                return await process_uploaded_file(UploadFile(file=io.BytesIO(b"source"), filename="note.txt"))

        with self.assertLogs("uvicorn.error", level="INFO") as captured:
            import asyncio
            result = asyncio.run(run_upload())

        self.assertEqual(result["_id"], "summary-id")
        self.assertTrue(any("Upload received" in message for message in captured.output))
        self.assertTrue(any("Content extraction finished" in message for message in captured.output))
        self.assertTrue(any("Local summarization finished" in message for message in captured.output))
        self.assertTrue(any("Upload processing completed" in message for message in captured.output))

    @patch(
        "services.summarizer_service.generate",
        return_value="Key Points\n- Starts May 3, from 6 pm to 8 pm.\nImportant Details\n- Weekdays.\nFinal Summary\nLibrary hours extend.",
    )
    def test_source_supported_dates_and_numbers_keep_generated_summary(self, _generate):
        source = "The library extends weekday hours from 6 pm to 8 pm starting on May 3."
        result = generate_structured_summary(source)
        self.assertIn("May 3", result)
        self.assertNotIn("unsupported claims", result)


if __name__ == "__main__":
    unittest.main()
