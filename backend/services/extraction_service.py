from io import BytesIO
from pathlib import Path
import wave

import av
import pdfplumber
from docx import Document
from PIL import Image, ImageOps, UnidentifiedImageError

from config import get_settings
from services.local_ai_service import transcribe_audio_bytes

SUPPORTED_EXTENSIONS = {
    ".txt": "txt",
    ".pdf": "pdf",
    ".docx": "docx",
    ".mp3": "audio",
    ".wav": "audio",
    ".mp4": "video",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".webp": "image",
    ".bmp": "image", ".tif": "image", ".tiff": "image",
}


def _detect_file_type(filename: str, content_type: str | None) -> str:
    extension = Path(filename).suffix.lower()
    if extension in SUPPORTED_EXTENSIONS:
        return SUPPORTED_EXTENSIONS[extension]

    if content_type:
        lowered = content_type.lower()
        if "text/plain" in lowered:
            return "txt"
        if "pdf" in lowered:
            return "pdf"
        if "word" in lowered or "officedocument" in lowered:
            return "docx"
        if lowered.startswith("audio/"):
            return "audio"
        if lowered.startswith("video/"):
            return "video"

    raise ValueError("Unsupported file type. Use PDF, DOCX, TXT, MP3, WAV, MP4, PNG, JPG, WEBP, BMP, or TIFF.")


def _extract_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode TXT file with supported encodings.")


def _extract_pdf(file_bytes: bytes) -> str:
    pages_text: list[str] = []
    rendered_pages: list[tuple[int, bytes]] = []
    visual_page_count = 0
    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        max_pages = get_settings().max_document_pages
        if len(pdf.pages) > max_pages:
            raise ValueError(f"PDF has {len(pdf.pages)} pages; the configured limit is {max_pages}.")
        for page_number, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text() or ""
            pages_text.append(page_text)
            has_visual_content = len(page_text.strip()) < 100 or bool(page.images)
            if has_visual_content:
                visual_page_count += 1
            if len(rendered_pages) < get_settings().max_vision_pages and has_visual_content:
                try:
                    import pymupdf
                    with pymupdf.open(stream=file_bytes, filetype="pdf") as document:
                        page_object = document[page_number - 1]
                        bounds = page_object.rect
                        scale = min(1.4, (get_settings().max_image_pixels / max(bounds.width * bounds.height, 1)) ** 0.5)
                        pixmap = page_object.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                        rendered_pages.append((page_number, pixmap.tobytes("png")))
                except Exception as exc:
                    raise RuntimeError(f"Could not render PDF page {page_number} for local visual analysis: {exc}") from exc
    if rendered_pages:
        from services.local_ai_service import generate
        visual_notes = []
        for page_number, image_bytes in rendered_pages:
            visual_notes.append(f"PDF page {page_number} visual analysis (do not repeat selectable text):\\n" + generate(
                "Describe visually meaningful charts, diagrams, tables, scans, and figures. Preserve visible labels, values, and units. Mark uncertain or unreadable information clearly.",
                image_bytes=image_bytes,
            ))
        pages_text.extend(visual_notes)
        if visual_page_count > get_settings().max_vision_pages:
            pages_text.append(f"Visual processing was limited to {get_settings().max_vision_pages} pages.")
    return "\n".join(pages_text)


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    content = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        content.extend(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells) for row in table.rows)
    return "\n".join(content)


def validate_image(file_bytes: bytes) -> bytes:
    try:
        with Image.open(BytesIO(file_bytes)) as image:
            if image.width * image.height > get_settings().max_image_pixels:
                raise ValueError(f"Image exceeds the configured {get_settings().max_image_pixels:,}-pixel limit.")
            image.verify()
        with Image.open(BytesIO(file_bytes)) as image:
            image = ImageOps.exif_transpose(image).convert("RGB")
            output = BytesIO()
            image.save(output, format="JPEG", quality=92)
            return output.getvalue()
    except ValueError:
        raise
    except (UnidentifiedImageError, Image.DecompressionBombError, OSError) as exc:
        raise ValueError("The uploaded image is invalid or uses an unsupported encoding.") from exc


def _extract_video_audio_bytes(file_bytes: bytes) -> bytes:
    pcm_parts: list[bytes] = []
    try:
        with av.open(BytesIO(file_bytes)) as container:
            audio_stream = next((stream for stream in container.streams if stream.type == "audio"), None)
            if audio_stream is None:
                raise ValueError("The video does not contain an audio track.")

            resampler = av.AudioResampler(format="s16", layout="mono", rate=16000)
            for frame in container.decode(audio_stream):
                pcm_parts.extend(
                    resampled.to_ndarray().tobytes()
                    for resampled in resampler.resample(frame)
                )
            pcm_parts.extend(
                resampled.to_ndarray().tobytes()
                for resampled in resampler.resample(None)
            )
    except av.error.FFmpegError as exc:
        raise ValueError("Could not decode the uploaded video or its audio track.") from exc

    pcm = b"".join(pcm_parts)
    if not pcm:
        raise ValueError("The video audio track is empty or could not be decoded.")

    output = BytesIO()
    with wave.open(output, "wb") as wav_file:
        wav_file.setnchannels(1)
        wav_file.setsampwidth(2)
        wav_file.setframerate(16000)
        wav_file.writeframes(pcm)
    return output.getvalue()


def _normalize_text(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return "\n".join(lines).strip()


def extract_text_from_file(filename: str, content_type: str | None, file_bytes: bytes) -> tuple[str, str, bytes | None]:
    if not filename:
        raise ValueError("File name is missing.")
    if not file_bytes:
        raise ValueError("Uploaded file is empty.")
    if len(file_bytes) > get_settings().max_upload_mb * 1024 * 1024:
        raise ValueError(f"File exceeds the {get_settings().max_upload_mb} MB upload limit.")

    file_type = _detect_file_type(filename, content_type)

    if file_type == "image":
        normalized_image = validate_image(file_bytes)
        return file_type, "Image uploaded for local visual analysis.", normalized_image
    if file_type == "txt":
        extracted = _extract_txt(file_bytes)
    elif file_type == "pdf":
        extracted = _extract_pdf(file_bytes)
    elif file_type == "docx":
        extracted = _extract_docx(file_bytes)
    elif file_type == "audio":
        extracted = transcribe_audio_bytes(file_bytes, filename)
    elif file_type == "video":
        audio_bytes = _extract_video_audio_bytes(file_bytes)
        audio_filename = f"{Path(filename).stem}.wav"
        extracted = transcribe_audio_bytes(audio_bytes, audio_filename)
    else:
        raise ValueError("Unsupported file type.")

    normalized = _normalize_text(extracted)
    if not normalized:
        raise ValueError("No readable text was extracted from the file.")

    return file_type, normalized, None
